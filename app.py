from flask import Flask, render_template, request, redirect, url_for, session, flash
from datetime import datetime
import random
import os
from openpyxl import Workbook, load_workbook

app = Flask(__name__)
app.secret_key = 'eclaro_vms_secret_key_2026'

EXCEL_FILE = 'visitor_logs.xlsx'

def init_excel():
    """ Gumawa ng Excel file kung wala pa """
    if not os.path.exists(EXCEL_FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = "Logs"
        # Table Headers
        ws.append(["Pass ID", "Full Name", "Contact Number", "Purpose", "Department", "Date & Time", "Status"])
        wb.save(EXCEL_FILE)

def append_to_excel(row_data):
    """ Mag-save ng bagong visitor sa Excel """
    init_excel()
    wb = load_workbook(EXCEL_FILE)
    ws = wb["Logs"]
    ws.append(row_data)
    wb.save(EXCEL_FILE)

def read_excel_logs():
    """ Magbasa ng logs mula sa Excel """
    init_excel()
    wb = load_workbook(EXCEL_FILE)
    ws = wb["Logs"]
    logs = []
    
    # Basahin mula row 2 (skip header)
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            logs.append({
                'id': row[0],
                'name': row[1],
                'phone': row[2],
                'purpose': row[3],
                'department': row[4],
                'time': row[5],
                'status': row[6]
            })
    return list(reversed(logs)) # Pinakabago sa itaas

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/register', methods=['POST'])
def register():
    name = request.form.get('fullName')
    phone = request.form.get('contactNumber')
    purpose = request.form.get('purpose')
    department = request.form.get('department')
    
    pass_id = f"ECL-{random.randint(100000, 999999)}"
    entry_time = datetime.now().strftime("%b %d, %Y - %I:%M %p")
    
    new_visitor = {
        'id': pass_id,
        'name': name,
        'phone': phone,
        'purpose': purpose,
        'department': department,
        'time': entry_time,
        'status': 'Inside'
    }
    
    # DIRECTLY SAVE TO EXCEL FILE (.xlsx)
    row = [pass_id, name, phone, purpose, department, entry_time, 'Inside']
    append_to_excel(row)
    
    return render_template('index.html', gate_pass=new_visitor)

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        username = request.form.get('username')
        password = request.form.get('password')
        
        if username == 'admin' and password == 'admin':
            session['logged_in'] = True
            return redirect(url_for('dashboard'))
        else:
            flash('Invalid username or password!', 'error')
            
    return render_template('login.html')

@app.route('/dashboard')
def dashboard():
    if not session.get('logged_in'):
        return redirect(url_for('login'))
        
    display_logs = read_excel_logs()
    
    total_registered = len(display_logs)
    currently_inside = sum(1 for v in display_logs if v['status'] == 'Inside')
    checked_out = sum(1 for v in display_logs if v['status'] == 'Checked Out')
    
    stats = {
        'total': total_registered,
        'inside': currently_inside,
        'out': checked_out
    }
    
    return render_template('dashboard.html', logs=display_logs, stats=stats)

@app.route('/toggle-status/<pass_id>')
def toggle_status(pass_id):
    if not session.get('logged_in'):
        return redirect(url_for('login'))
        
    init_excel()
    wb = load_workbook(EXCEL_FILE)
    ws = wb["Logs"]
    
    for row in ws.iter_rows(min_row=2):
        if row[0].value == pass_id:
            current_status = row[6].value
            row[6].value = 'Checked Out' if current_status == 'Inside' else 'Inside'
            break
            
    wb.save(EXCEL_FILE)
    return redirect(url_for('dashboard'))

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))

if __name__ == '__main__':
    app.run(debug=True)
